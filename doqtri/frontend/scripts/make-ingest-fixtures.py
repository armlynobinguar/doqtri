"""
Regenerates the ingest regression fixtures in e2e/fixtures/ingest/.

    pip install python-docx python-pptx
    python scripts/make-ingest-fixtures.py

The DOCX files cover the table shapes that used to hang or garble ingest:
a plain table, merged cells, a table nested in a cell, and a long table. The
content is synthetic so the files are safe to commit.
"""

from pathlib import Path

import docx
import pptx
from pptx.util import Inches

OUT = Path(__file__).resolve().parent.parent / "e2e" / "fixtures" / "ingest"
OUT.mkdir(parents=True, exist_ok=True)


def table(document, rows, cols, fill):
    t = document.add_table(rows=rows, cols=cols)
    t.style = "Table Grid"
    for r in range(rows):
        for c in range(cols):
            t.cell(r, c).text = fill(r, c)
    return t


def docx_simple():
    d = docx.Document()
    d.add_heading("Sprint Plan", 1)
    d.add_paragraph("The plan for the ingest sprint, with owners and dates.")
    d.add_heading("Milestones", 2)
    rows = [("Milestone", "Owner", "Due"), ("Streaming ingest", "Dev", "Week 1"),
            ("Wallet session", "Dev", "Week 2"), ("Audit page", "Dev", "Week 3")]
    table(d, len(rows), 3, lambda r, c: rows[r][c])
    d.add_heading("Risks", 2)
    d.add_paragraph("Tables with | pipes | in cells must be escaped.")
    d.save(OUT / "docx-table-simple.docx")


def docx_merged():
    d = docx.Document()
    d.add_heading("Budget", 1)
    t = table(d, 4, 4, lambda r, c: f"r{r}c{c}")
    t.cell(0, 0).merge(t.cell(0, 2))  # colspan 3
    t.cell(1, 1).merge(t.cell(3, 1))  # rowspan 3
    d.add_paragraph("After the merged table.")
    d.save(OUT / "docx-table-merged.docx")


def docx_nested():
    d = docx.Document()
    d.add_heading("Nested", 1)
    t = table(d, 2, 2, lambda r, c: "")
    t.cell(0, 0).text = "outer"
    inner = t.cell(1, 1).add_table(rows=2, cols=2)
    for r in range(2):
        for c in range(2):
            inner.cell(r, c).text = f"in{r}{c}"
    d.save(OUT / "docx-table-nested.docx")


def docx_large():
    d = docx.Document()
    d.add_heading("Transaction Log", 1)
    d.add_paragraph("Four hundred rows: the shape that outran the function before.")
    table(d, 400, 6, lambda r, c: "Column " + str(c) if r == 0 else f"row {r} value {c}")
    d.save(OUT / "docx-table-large.docx")


def pptx_deck():
    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Doqtri Roadmap"
    slide.placeholders[1].text = "Reliable ingest\nStable wallet sessions\nPublic audit page"
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Targets"
    shape = slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(1.5))
    for r, row in enumerate([("Metric", "Target"), ("Testers", "20"), ("Transactions", "60")]):
        for c, value in enumerate(row):
            shape.table.cell(r, c).text = value
    deck.save(OUT / "pptx-deck.pptx")


def pdf_minimal():
    """A one-page text PDF written by hand, so no PDF library is needed."""
    text = "Doqtri PDF fixture. Heading: Scope. Body: the ingest pipeline."
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    (OUT / "pdf-minimal.pdf").write_bytes(bytes(out))


def text_notes():
    (OUT / "text-notes.txt").write_text(
        "﻿Meeting notes\n\nDecisions:\n- Ship streaming ingest first\n- Audit page in week 3\n",
        encoding="utf-8",
    )


for make in (docx_simple, docx_merged, docx_nested, docx_large, pptx_deck, pdf_minimal, text_notes):
    make()
print(f"wrote fixtures to {OUT}")
